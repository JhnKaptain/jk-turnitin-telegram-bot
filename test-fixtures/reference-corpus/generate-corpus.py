import json, sys, os, random
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = sys.argv[1]
random.seed(3)

BODY = [
 "Osteoarthritis is a degenerative joint condition that leads to the gradual breakdown of articular cartilage and remodelling of the underlying bone. The patient described morning stiffness lasting less than thirty minutes, which is consistent with a mechanical rather than inflammatory process (Brophy & Fillingham, 2022).",
 "Management should begin with education and a structured exercise programme, because strengthening the quadriceps reduces pain and improves function for most patients. Weight reduction is also recommended when the body mass index exceeds the healthy range, and the clinician should revisit these goals at every follow-up visit.",
 "Pharmacological treatment is reserved for symptoms that do not respond to conservative measures. Topical anti-inflammatory agents are preferred before oral agents in older adults, since gastrointestinal and renal risks increase with age and with concurrent medication use.",
 "Screening recommendations must be reviewed at the same consultation. According to current guidance, adults aged forty-five to seventy-five should be offered colorectal cancer screening, and women aged forty to seventy-four should be offered biennial mammography unless a higher-risk pathway applies.",
 "In summary, the evidence supports a stepped approach that prioritises non-pharmacological care, reserves medication for persistent symptoms, and integrates preventive screening into routine visits so that the patient's long-term health is protected.",
]
AFTER_PROSE = "The appendix presents the interview schedule that was used during the placement, and it explains how each question relates to the learning outcomes described in the reflective sections above. Participants were given the questions in advance so that they could prepare their answers carefully."

APA = [
 "Brophy, R. H., & Fillingham, Y. A. (2022). AAOS clinical practice guideline summary: Management of osteoarthritis of the knee (nonarthroplasty). Journal of the American Academy of Orthopaedic Surgeons, 30(9), e721–e729. https://doi.org/10.5435/JAAOS-D-21-01233",
 "Jurado-Castro, J. M., Munoz-Lopez, M., Ledesma, A. S. T., & Ranchal-Sanchez, A. (2022). Effectiveness of exercise in patients with overweight or obesity suffering from knee osteoarthritis: A systematic review and meta-analysis. International Journal of Environmental Research and Public Health, 19(17), 10510. https://doi.org/10.3390/ijerph191710510",
 "Moseng, T., Vliet Vlieland, T. P., Battista, S., Beckwée, D., Boyadzhieva, V., Conaghan, P. G., & Østerås, N. (2024). EULAR recommendations for the non-pharmacological core management of hip and knee osteoarthritis: 2023 update. Annals of the Rheumatic Diseases, 83(6), 730–740.",
 "U.S. Preventive Services Task Force. (2021). Screening for colorectal cancer: US Preventive Services Task Force recommendation statement. JAMA, 325(19), 1965–1977. https://doi.org/10.1001/jama.2021.6238",
 "Cowley, S., & McCarthy, G. (2023). Diagnosis and treatment of calcium pyrophosphate deposition (CPPD) disease: A review. Open Access Rheumatology: Research and Reviews, 15, 33–41.",
 "Smith, M. H., & Berman, J. R. (2022). What is rheumatoid arthritis? JAMA, 327(12), 1194. https://doi.org/10.1001/jama.2022.0786",
 "Mohseni, M., Mabrouk, A., & Li, D. (2024). Pes anserine bursitis. In StatPearls [Internet]. StatPearls Publishing. https://www.ncbi.nlm.nih.gov/books/NBK532941/",
 "World Health Organization. (n.d.). Musculoskeletal health. Retrieved March 3, 2025, from https://www.who.int/news-room/fact-sheets/detail/musculoskeletal-conditions",
 "Poznyak, A. V., Kirichenko, T. V., Beloyartsev, D. F., Churov, A. V., Kovyanova, T. I., Starodubtseva, I. A., & Orekhov, A. N. (2024). Rheumatoid arthritis: What inflammation do we face? Journal of Molecular Pathology, 5(4), 454–465.",
 "Porter, L. (2016). Young children's behaviour: Guidance approaches for early childhood educators (4th ed.). Allen & Unwin.",
]
HARVARD = [
 "Brophy, R.H. and Fillingham, Y.A. (2022) 'AAOS clinical practice guideline summary: management of osteoarthritis of the knee', Journal of the American Academy of Orthopaedic Surgeons, 30(9), pp. e721–e729. doi:10.5435/JAAOS-D-21-01233.",
 "Greggs plc (2026) Annual report and accounts 2025. Newcastle upon Tyne: Greggs plc.",
 "National Institute for Health and Care Excellence (2022) Osteoarthritis in over 16s: diagnosis and management. NICE guideline NG226. Available at: https://www.nice.org.uk/guidance/ng226 (Accessed: 3 March 2025).",
 "Moseng, T. et al. (2024) 'EULAR recommendations for the non-pharmacological core management of hip and knee osteoarthritis: 2023 update', Annals of the Rheumatic Diseases, 83(6), pp. 730–740.",
 "Porter, L. (2016) Young children's behaviour: guidance approaches for early childhood educators. 4th edn. Sydney: Allen & Unwin.",
 "World Health Organization (2021) Musculoskeletal health. Geneva: WHO. Available at: https://www.who.int/news-room/fact-sheets/detail/musculoskeletal-conditions (Accessed: 12 May 2025).",
 "Smith, M.H. and Berman, J.R. (2022) 'What is rheumatoid arthritis?', JAMA, 327(12), p. 1194.",
]
CHICAGO_AD = [
 "Brophy, Robert H., and Yale A. Fillingham. 2022. \"AAOS Clinical Practice Guideline Summary: Management of Osteoarthritis of the Knee.\" Journal of the American Academy of Orthopaedic Surgeons 30 (9): e721–29. https://doi.org/10.5435/JAAOS-D-21-01233.",
 "Elliott, Anthony. 2010. The Routledge Companion to Social Theory. London: Routledge.",
 "Moseng, Tuva, Thea Vliet Vlieland, and Nina Østerås. 2024. \"EULAR Recommendations for the Non-Pharmacological Core Management of Hip and Knee Osteoarthritis: 2023 Update.\" Annals of the Rheumatic Diseases 83 (6): 730–40.",
 "Porter, Louise. 2016. Young Children's Behaviour: Guidance Approaches for Early Childhood Educators. 4th ed. Sydney: Allen & Unwin.",
 "World Health Organization. 2021. \"Musculoskeletal Health.\" Fact sheet. Geneva: WHO. https://www.who.int/news-room/fact-sheets/detail/musculoskeletal-conditions.",
 "Smith, Melissa H., and Jessica R. Berman. 2022. \"What Is Rheumatoid Arthritis?\" JAMA 327 (12): 1194.",
]
CHICAGO_NB = [
 "Brophy, Robert H., and Yale A. Fillingham. \"AAOS Clinical Practice Guideline Summary: Management of Osteoarthritis of the Knee.\" Journal of the American Academy of Orthopaedic Surgeons 30, no. 9 (2022): e721–29.",
 "Elliott, Anthony. The Routledge Companion to Social Theory. London: Routledge, 2010.",
 "Moseng, Tuva, Thea Vliet Vlieland, and Nina Østerås. \"EULAR Recommendations for the Non-Pharmacological Core Management of Hip and Knee Osteoarthritis.\" Annals of the Rheumatic Diseases 83, no. 6 (2024): 730–40.",
 "Porter, Louise. Young Children's Behaviour: Guidance Approaches for Early Childhood Educators. 4th ed. Sydney: Allen & Unwin, 2016.",
 "Smith, Melissa H., and Jessica R. Berman. \"What Is Rheumatoid Arthritis?\" JAMA 327, no. 12 (2022): 1194.",
 "World Health Organization. Musculoskeletal Health. Geneva: WHO, 2021. https://www.who.int/news-room/fact-sheets/detail/musculoskeletal-conditions.",
]
IEEE = [
 "[1] R. H. Brophy and Y. A. Fillingham, \"AAOS clinical practice guideline summary: Management of osteoarthritis of the knee,\" J. Amer. Acad. Orthop. Surg., vol. 30, no. 9, pp. e721–e729, 2022, doi: 10.5435/JAAOS-D-21-01233.",
 "[2] T. Moseng, T. P. Vliet Vlieland, and N. Østerås, \"EULAR recommendations for the non-pharmacological core management of hip and knee osteoarthritis: 2023 update,\" Ann. Rheum. Dis., vol. 83, no. 6, pp. 730–740, 2024.",
 "[3] A. Kumar and B. Chen, \"Edge analytics for industrial sensors,\" IEEE Trans. Ind. Informat., vol. 9, no. 2, pp. 1–8, Feb. 2021.",
 "[4] World Health Organization, \"Musculoskeletal health,\" 2021. [Online]. Available: https://www.who.int/news-room/fact-sheets/detail/musculoskeletal-conditions",
 "[5] M. H. Smith and J. R. Berman, \"What is rheumatoid arthritis?\" JAMA, vol. 327, no. 12, p. 1194, 2022.",
 "[6] L. Porter, Young Children's Behaviour, 4th ed. Sydney, Australia: Allen & Unwin, 2016.",
]
VANCOUVER = [
 "1. Brophy RH, Fillingham YA. AAOS clinical practice guideline summary: management of osteoarthritis of the knee (nonarthroplasty). J Am Acad Orthop Surg. 2022;30(9):e721-e729.",
 "2. Jurado-Castro JM, Munoz-Lopez M, Ledesma AST, Ranchal-Sanchez A. Effectiveness of exercise in patients with overweight or obesity suffering from knee osteoarthritis: a systematic review and meta-analysis. Int J Environ Res Public Health. 2022;19(17):10510.",
 "3. Moseng T, Vliet Vlieland TP, Battista S, et al. EULAR recommendations for the non-pharmacological core management of hip and knee osteoarthritis: 2023 update. Ann Rheum Dis. 2024;83(6):730-740.",
 "4. US Preventive Services Task Force. Screening for colorectal cancer: US Preventive Services Task Force recommendation statement. JAMA. 2021;325(19):1965-1977.",
 "5. Smith MH, Berman JR. What is rheumatoid arthritis? JAMA. 2022;327(12):1194.",
 "6. World Health Organization. Musculoskeletal health [Internet]. Geneva: WHO; 2021 [cited 2025 Mar 3]. Available from: https://www.who.int/news-room/fact-sheets/detail/musculoskeletal-conditions",
]
MLA = [
 "Brophy, Robert H., and Yale A. Fillingham. \"AAOS Clinical Practice Guideline Summary: Management of Osteoarthritis of the Knee.\" Journal of the American Academy of Orthopaedic Surgeons, vol. 30, no. 9, 2022, pp. e721–29.",
 "Elliott, Anthony. The Routledge Companion to Social Theory. Routledge, 2010.",
 "Moseng, Tuva, et al. \"EULAR Recommendations for the Non-Pharmacological Core Management of Hip and Knee Osteoarthritis: 2023 Update.\" Annals of the Rheumatic Diseases, vol. 83, no. 6, 2024, pp. 730–40.",
 "Porter, Louise. Young Children's Behaviour: Guidance Approaches for Early Childhood Educators. 4th ed., Allen & Unwin, 2016.",
 "\"Musculoskeletal Health.\" World Health Organization, 2021, www.who.int/news-room/fact-sheets/detail/musculoskeletal-conditions. Accessed 3 Mar. 2025.",
 "Smith, Melissa H., and Jessica R. Berman. \"What Is Rheumatoid Arthritis?\" JAMA, vol. 327, no. 12, 2022, p. 1194.",
]
OSCOLA = [
 "Donoghue v Stevenson [1932] AC 562 (HL).",
 "Caparo Industries plc v Dickman [1990] 2 AC 605.",
 "Montgomery v Lanarkshire Health Board [2015] UKSC 11, [2015] AC 1430.",
 "Human Rights Act 1998, s 6.",
 "Mental Capacity Act 2005.",
 "Jonathan Herring, Medical Law and Ethics (9th edn, OUP 2022).",
 "Emily Jackson, 'Informed consent to medical treatment and the impotence of tort' in Sheila McLean (ed), First Do No Harm (Ashgate 2006).",
 "Jose Miola, 'On the materiality of risk: paper tigers and panaceas' (2009) 17 Medical Law Review 76.",
]
URL_ONLY = [
 "https://www.who.int/news-room/fact-sheets/detail/musculoskeletal-conditions",
 "https://www.nice.org.uk/guidance/ng226",
 "https://www.uspreventiveservicestaskforce.org/uspstf/recommendation/colorectal-cancer-screening",
 "https://www.ncbi.nlm.nih.gov/books/NBK532941/",
 "https://doi.org/10.1001/jama.2022.0786",
]
STYLES = {"apa":APA,"harvard":HARVARD,"chicago_ad":CHICAGO_AD,"chicago_nb":CHICAGO_NB,"ieee":IEEE,"vancouver":VANCOUVER,"mla":MLA,"oscola":OSCOLA,"url_only":URL_ONLY,
          "mixed":[APA[0],URL_ONLY[0],HARVARD[2],APA[5],URL_ONLY[3],CHICAGO_NB[1]]}

def add_page_field(paragraph):
    run = paragraph.add_run()
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = 'PAGE'
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'end')
    run._r.append(f1); run._r.append(it); run._r.append(f2)

def body(d, n=5, mid_heading=None, mention=False, toc=False, cover=False):
    if cover:
        for t in ["Case Study: Evaluating and Managing Cardiovascular Conditions","Student Name","Course Code: NURS 6512","Date: 14 October 2026"]:
            p=d.add_paragraph(t); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        d.add_page_break()
    if toc:
        d.add_heading("Table of Contents",1)
        for t,n2 in [("Introduction","1"),("Assessment","2"),("Management Plan","3"),("References","5"),("Appendix A","6")]:
            p=d.add_paragraph(); p.paragraph_format.tab_stops.add_tab_stop(Inches(6)); p.add_run(t+"\t"+n2)
        d.add_page_break()
    d.add_heading("Introduction",1)
    for i in range(n):
        if mid_heading and i==2: d.add_heading(mid_heading,2)
        t=BODY[i%len(BODY)]
        if mention and i==1: t += " Full details of every source are given in the References section at the end of this paper."
        p=d.add_paragraph(t); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY

def refs(d, entries, heading="References", heading_style="heading", fmt="hanging", blank_between=False,
         per_entry_font=None, flatten=False, letterspace=False, smallcaps=False, caps=False, trailing_period=False,
         in_table=False, as_list=False, stray_quote_index=None, page_break_before=False, spacing=None):
    if page_break_before: d.add_page_break()
    htext = heading
    if letterspace: htext = " ".join(heading)
    if caps: htext = htext.upper()
    if trailing_period: htext += "."
    if heading_style=="heading":
        h=d.add_heading(htext,1)
    elif heading_style=="bold":
        h=d.add_paragraph(); r=h.add_run(htext); r.bold=True
    elif heading_style=="centered_bold":
        h=d.add_paragraph(); r=h.add_run(htext); r.bold=True; h.alignment=WD_ALIGN_PARAGRAPH.CENTER
    elif heading_style=="plain":
        h=d.add_paragraph(htext)
    elif heading_style=="underline":
        h=d.add_paragraph(); r=h.add_run(htext); r.underline=True; r.bold=True
    if smallcaps:
        for r in h.runs: r.font.small_caps=True
    if flatten:
        # heading and first entry in the same paragraph, separated by a manual line break
        h.runs[-1].add_break(WD_BREAK.LINE)
        h.add_run(entries[0]).bold=False
        for r in h.runs[1:]: r.bold=False
        entries=entries[1:]
    if in_table:
        t=d.add_table(rows=len(entries),cols=1); t.style="Table Grid"
        for i,e in enumerate(entries): t.cell(i,0).text=e
        return
    for i,e in enumerate(entries):
        if stray_quote_index==i: e="‘"+e
        p=d.add_paragraph(e, style="List Number" if as_list else None)
        pf=p.paragraph_format
        if fmt=="hanging": pf.left_indent=Inches(0.5); pf.first_line_indent=Inches(-0.5)
        if spacing: pf.line_spacing=spacing
        if per_entry_font and i in per_entry_font:
            for r in p.runs: r.font.size=Pt(per_entry_font[i])
        if blank_between: d.add_paragraph("")

def after(d, kind):
    if kind=="none": return None
    if kind=="appendix":
        d.add_heading("Appendix A: Interview Schedule",1); d.add_paragraph(AFTER_PROSE); return AFTER_PROSE[:60]
    if kind=="appendix_numbered":
        d.add_heading("Appendix A",1)
        for q in ["How do you decide when to refer a patient for imaging?","Which non-pharmacological options do you recommend first?","How often do you review screening status during visits?"]:
            d.add_paragraph(q, style="List Number")
        return "How do you decide when to refer a patient for imaging?"
    if kind=="declaration":
        h=d.add_paragraph(); r=h.add_run("Generative AI Use Declaration: Mandatory"); r.bold=True
        d.add_paragraph("I declare that generative AI tools were not used to produce the content of this assessment, and that all sources consulted have been acknowledged in the reference list above.")
        return "I declare that generative AI tools were not used"
    if kind=="task2":
        d.add_page_break(); d.add_heading("Task 2: Reflective Account",1)
        for t in BODY[:3]: d.add_paragraph(t)
        return BODY[0][:60]
    if kind=="second_list":
        d.add_page_break(); d.add_heading("Task 2: Reflective Account",1)
        for t in BODY[:3]: d.add_paragraph(t)
        refs(d, HARVARD[:4], heading="Reference List")
        return BODY[0][:60]
    if kind=="wordcount":
        d.add_paragraph("Word count: 1,498 (excluding references)"); return None
    if kind=="table":
        d.add_heading("Appendix B: Screening Summary",1)
        t=d.add_table(rows=3,cols=2); t.style="Table Grid"
        for r,row in enumerate([("Screening","Interval"),("Colorectal cancer","Every 10 years"),("Mammography","Every 2 years")]):
            for c,v in enumerate(row): t.cell(r,c).text=v
        return None

specs=[]
def make(name, style="apa", entries=None, heading="References", after_kind="none", header=None, footer_page=False, body_kwargs=None, **rk):
    d=Document()
    sec=d.sections[0]
    if header:
        hp=sec.header.paragraphs[0]; hp.text=header; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        if header.endswith(" "): add_page_field(hp)
    if footer_page:
        fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; add_page_field(fp)
    body(d, **(body_kwargs or {}))
    ents=entries if entries is not None else STYLES[style]
    refs(d, ents, heading=heading, **rk)
    keep_after=after(d, after_kind)
    path=os.path.join(OUT, name+".docx"); d.save(path)
    specs.append({"name":name,"style":style,"heading":heading,"after":after_kind,
                  "entries":[e.lstrip("‘")[:28] for e in ents],"keepBefore":BODY[0][:60],"keepAfter":keep_after,
                  "secondList":after_kind=="second_list"})

# 1. every style, standard formatting
for s in STYLES: make(f"style_{s}", style=s)
# 2. heading variants (APA)
for h in ["REFERENCES","Reference List","Reference list:","Bibliography","Works Cited","Sources","Literature Cited","Citations",
          "References and Bibliography","6. References","6.0 References","Chapter 6: References","Reference","List of References",
          "Works Consulted","Sources Cited","Reference Sources","Further Reading","Bibliography and References","References Cited"]:
    make("heading_"+h.lower().replace(" ","_").replace(":","").replace(".","_"), heading=h)
# 3. heading formatting
make("hfmt_bold", heading_style="bold"); make("hfmt_centered_bold", heading_style="centered_bold"); make("hfmt_plain", heading_style="plain")
make("hfmt_underline", heading_style="underline"); make("hfmt_smallcaps", heading_style="bold", smallcaps=True, heading="List of References")
make("hfmt_letterspaced", heading_style="centered_bold", letterspace=True); make("hfmt_caps_period", caps=True, trailing_period=True)
make("hfmt_flattened", flatten=True, heading_style="bold"); make("hfmt_pagebreak_before", page_break_before=True)
# 4. entry formatting quirks
make("efmt_no_indent", fmt="none"); make("efmt_double_spaced", spacing=2.0); make("efmt_blank_between", blank_between=True, fmt="none")
make("efmt_mixed_fonts", per_entry_font={2:10,4:9,6:13}, fmt="none"); make("efmt_stray_quote_noindent", stray_quote_index=4, fmt="none", per_entry_font={4:10})
make("efmt_in_table", in_table=True); make("efmt_auto_numbered", as_list=True, fmt="none"); make("efmt_two_entries", entries=APA[:2])
make("efmt_long_list", entries=APA+HARVARD+CHICAGO_AD+MLA)
make("efmt_ieee_double_spaced", style="ieee", spacing=2.0, fmt="none"); make("efmt_vancouver_noindent", style="vancouver", fmt="none")
make("efmt_chicago_nb_double", style="chicago_nb", spacing=2.0)
# 5. endings
for k in ["appendix","appendix_numbered","declaration","task2","second_list","wordcount","table"]: make("after_"+k, after_kind=k)
make("after_appendix_harvard", style="harvard", after_kind="appendix"); make("after_task2_ieee", style="ieee", after_kind="task2")
# 6. running headers / footers / body quirks
make("run_mla_header", style="mla", header="Smith ", footer_page=False); make("run_footer_page", footer_page=True, entries=APA+HARVARD)
make("run_title_header", header="CASE STUDY: CARDIOVASCULAR CONDITIONS", footer_page=True, entries=APA+HARVARD)
make("body_mention_and_toc", body_kwargs={"mention":True,"toc":True}); make("body_mid_heading", body_kwargs={"mid_heading":"Reference Framework for Assessment"})
make("body_cover_page", body_kwargs={"cover":True}, after_kind="appendix")
json.dump(specs, open(os.path.join(OUT,"specs.json"),"w"), indent=1)
print(len(specs),"documents")

# ---- negative / safety documents ----
def make_neg(name, build, expectReview=False, keepAfter=None, entries=None, noRefs=False):
    d=Document(); build(d); path=os.path.join(OUT,name+".docx"); d.save(path)
    specs.append({"name":name,"style":"neg","heading":"","after":"","entries":entries or [], "keepBefore":BODY[0][:60],
                  "keepAfter":keepAfter,"secondList":False,"expectReview":expectReview,"noRefs":noRefs})
CITED_PROSE="Brophy and Fillingham (2022) argue that conservative management should precede surgery, while Moseng et al. (2024) emphasise exercise. Smith and Berman (2022) describe the inflammatory pathway in detail (see https://doi.org/10.1001/jama.2022.0786 for the full statement)."
def b_sources_of_error(d):
    body(d); d.add_heading("Sources of Error",2)
    d.add_paragraph("Several sources of error affected this evaluation. The sample was small, follow-up was short, and the outcome measures were self-reported, which may have inflated the apparent benefit of the exercise programme described above.")
    d.add_heading("Further Reading",2); d.add_paragraph("Readers who want more depth should consult the guideline summaries listed below and the systematic reviews published since 2022, which extend the evidence base considerably.")
    refs(d, APA)
make_neg("neg_sources_of_error", b_sources_of_error, entries=[e[:28] for e in APA], keepAfter="Several sources of error affected this evaluation")
def b_no_refs(d):
    body(d); d.add_paragraph(CITED_PROSE); d.add_paragraph(CITED_PROSE)
make_neg("neg_no_references", b_no_refs, noRefs=True, keepAfter=CITED_PROSE[:50])
def b_intext_starts(d):
    body(d, n=2)
    for _ in range(3): d.add_paragraph(CITED_PROSE)
    refs(d, APA)
make_neg("neg_intext_sentence_starts", b_intext_starts, entries=[e[:28] for e in APA], keepAfter=CITED_PROSE[:50])
def b_appendix_cited(d):
    body(d); refs(d, APA); d.add_heading("Appendix A: Extended Discussion",1)
    d.add_paragraph(CITED_PROSE); d.add_paragraph(AFTER_PROSE)
make_neg("neg_appendix_with_citations", b_appendix_cited, entries=[e[:28] for e in APA], keepAfter=CITED_PROSE[:50])
def b_prose_between(d):
    body(d); refs(d, APA[:3]); d.add_paragraph("In my clinical placement I observed that nurses explained every medication to the patient before giving it, which reduced anxiety considerably and improved adherence over the following weeks."); refs(d, APA[3:6], heading=None) if False else None
    for e in APA[3:6]:
        p=d.add_paragraph(e); p.paragraph_format.left_indent=Inches(0.5); p.paragraph_format.first_line_indent=Inches(-0.5)
make_neg("neg_prose_between_entries", b_prose_between, expectReview=True)
def b_unheaded_tail(d):
    body(d); refs(d, APA)
    d.add_paragraph("Reflecting on this case, I realise that my initial assessment relied too heavily on imaging and not enough on the patient's own account of her symptoms, which is something I intend to change in future consultations.")
make_neg("neg_unheaded_prose_after_refs", b_unheaded_tail, expectReview=True)
def b_oscola_table_cases(d):
    body(d); refs(d, OSCOLA[:3], heading="Table of Cases"); refs(d, OSCOLA[3:5], heading="Table of Legislation"); refs(d, OSCOLA[5:], heading="Bibliography")
make_neg("neg_oscola_three_tables", b_oscola_table_cases, entries=[e[:28] for e in OSCOLA])
json.dump(specs, open(os.path.join(OUT,"../specs.json"),"w"), indent=1)
print(len(specs),"documents (with negatives)")
